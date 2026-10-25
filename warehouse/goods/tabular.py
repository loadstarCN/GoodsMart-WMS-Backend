"""商品导入导出的表格文件：CSV（UTF-8 / Shift_JIS）与 Excel（.xlsx，openpyxl）的读与写。

读：
- CSV：编码 auto（默认）= 有 BOM 按 UTF-8；内容是合法 UTF-8 按 UTF-8（Shift_JIS 的日文几乎不会恰好是合法 UTF-8）；
  否则按 Shift_JIS（CP932）。也可手动指定 utf-8 / shift_jis（指定了就只按它解）
- Excel：第一个工作表（或 sheet 指定的）；读缓存值（公式取上次计算的结果）；单元格统一转成文字再交给导入校验
  （整数 → 140、小数 → 0.328、日期 → 2025-07-01）。打开前检查解压后的大小，防压缩炸弹
- 行号 = 文件里的行号（表头所在行为第 1 行时与 Excel / 文本编辑器一致）；整行为空的行跳过（行号照算）
写：
- CSV：UTF-8 带 BOM，或 Shift_JIS（CP932）。CP932 先把 JIS 与 Unicode 映射不一致的字（〜 ‖ − ¢ £ ¬ —）换成 CP932 里的对应字，
  仍编码不了的字符写成「?」（unencodable_chars 预先统计）
- Excel：write_only 模式逐行写；以 = 开头的文字按文字写（不当公式）；去掉 Excel 不允许的控制字符；超过单元格上限（32,767 字）截断
"""
import csv
import io
import re
import zipfile
from datetime import date, datetime, time
from decimal import Decimal

ENCODING_LABELS = {'utf-8': 'UTF-8', 'utf-8-sig': 'UTF-8 (BOM)', 'cp932': 'Shift_JIS (CP932)'}
_ENCODING_ALIASES = {
    'auto': 'auto', '': 'auto',
    'utf-8': 'utf-8', 'utf8': 'utf-8', 'utf-8-sig': 'utf-8', 'utf8-sig': 'utf-8', 'utf-8-bom': 'utf-8',
    'shift_jis': 'cp932', 'shift-jis': 'cp932', 'shiftjis': 'cp932', 'sjis': 'cp932', 'cp932': 'cp932',
    'windows-31j': 'cp932', 'ms932': 'cp932',
}
FORMATS = ('csv', 'xlsx')
CONTENT_TYPES = {
    'csv': 'text/csv',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
}
# 读 Excel 的上限：解压后的总大小、每行读的列数
XLSX_MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_COLUMNS = 200
EXCEL_CELL_MAX_CHARS = 32767

# JIS X 0208 与 Unicode 的映射在 CP932 里另有对应字：先换过去再编码（否则 Mac / Web 上输入的「〜」会变成「?」）
CP932_SUBSTITUTES = str.maketrans({
    '〜': '～',   # 〜 WAVE DASH → ～
    '‖': '∥',   # ‖ → ∥
    '−': '－',   # − MINUS SIGN → －
    '¢': '￠',   # ¢ → ￠
    '£': '￡',   # £ → ￡
    '¬': '￢',   # ¬ → ￢
    '—': '―',   # — EM DASH → ―
})
_ILLEGAL_XML_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


class TableError(ValueError):
    """文件读不了：reason = NOT_XLSX / NOT_CSV / UNKNOWN_ENCODING / NOT_UTF8 / NOT_SHIFT_JIS / EMPTY / SHEET_NOT_FOUND /
    TOO_MANY_ROWS / TOO_LARGE；details 附加信息"""

    def __init__(self, reason, message, **details):
        super().__init__(message)
        self.reason = reason
        self.message = message
        self.details = details


def normalize_encoding(value):
    """编码参数 → auto / utf-8 / cp932；不认识返回 None"""
    if value is None:
        return 'auto'
    return _ENCODING_ALIASES.get(str(value).strip().lower())


def file_format(filename):
    """文件名 → csv / xlsx；别的返回 None"""
    name = (filename or '').strip().lower()
    if name.endswith('.csv'):
        return 'csv'
    if name.endswith('.xlsx'):
        return 'xlsx'
    return None


class Table:
    """读出来的表：rows = [(行号, [单元格文字, …])]（整行为空的不在里面）"""

    def __init__(self, fmt, rows, encoding=None, encoding_detected=False, sheet=None, sheets=None):
        self.format = fmt
        self.rows = rows
        self.encoding = encoding
        self.encoding_detected = encoding_detected
        self.sheet = sheet
        self.sheets = sheets or []


def _decode(content: bytes, encoding: str):
    """→ (文字, 实际编码 utf-8-sig / utf-8 / cp932, 是否自动识别)"""
    if encoding == 'auto':
        if content.startswith(b'\xef\xbb\xbf'):
            try:
                return content[3:].decode('utf-8'), 'utf-8-sig', True
            except UnicodeDecodeError:
                raise TableError('UNKNOWN_ENCODING', "The file starts with a UTF-8 BOM but is not valid UTF-8")
        try:
            return content.decode('utf-8'), 'utf-8', True
        except UnicodeDecodeError:
            pass
        try:
            return content.decode('cp932'), 'cp932', True
        except UnicodeDecodeError:
            raise TableError('UNKNOWN_ENCODING', "The CSV file is neither UTF-8 nor Shift_JIS text")
    if encoding == 'utf-8':
        try:
            if content.startswith(b'\xef\xbb\xbf'):
                return content[3:].decode('utf-8'), 'utf-8-sig', False
            return content.decode('utf-8'), 'utf-8', False
        except UnicodeDecodeError:
            raise TableError('NOT_UTF8', "The CSV file must be UTF-8 encoded text.")
    try:
        return content.decode('cp932'), 'cp932', False
    except UnicodeDecodeError:
        raise TableError('NOT_SHIFT_JIS', "The CSV file is not Shift_JIS text.")


def _trim(cells):
    """去掉行尾的空单元格（Excel 的只读模式按最大列数补 None）；整行为空 → []"""
    end = len(cells)
    while end and not str(cells[end - 1]).strip():
        end -= 1
    return list(cells[:end])


def _limit_rows(rows, max_rows):
    """数据行（表头之后的非空行）不能超过 max_rows"""
    if max_rows is not None and len(rows) - 1 > max_rows:
        raise TableError('TOO_MANY_ROWS', f"The file has more than {max_rows} data rows", max_rows=max_rows)


def read_csv(content: bytes, encoding='auto', max_rows=None) -> Table:
    text, used, detected = _decode(content, encoding)
    reader = csv.reader(io.StringIO(text, newline=''))
    rows = []
    try:
        for number, cells in enumerate(reader, start=1):
            cells = _trim(cells[:MAX_COLUMNS])
            if cells:
                rows.append((number, cells))
                if max_rows is not None and len(rows) - 1 > max_rows:
                    break
    except csv.Error as exc:
        raise TableError('NOT_CSV', f"The file is not a valid CSV file: {exc}")
    _limit_rows(rows, max_rows)
    return Table('csv', rows, encoding=used, encoding_detected=detected)


def cell_text(value) -> str:
    """Excel 单元格 → 文字（导入校验统一按文字处理）"""
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return '%.15g' % value
    if isinstance(value, Decimal):
        return format(value, 'f')
    if isinstance(value, datetime):
        if value.time() == time():
            return value.date().isoformat()
        return value.isoformat(sep=' ', timespec='seconds')
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _check_zip(content: bytes):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        raise TableError('NOT_XLSX', "The file is not a valid Excel (.xlsx) file")
    if total > XLSX_MAX_UNCOMPRESSED_BYTES:
        raise TableError('TOO_LARGE', "The Excel file is too large when uncompressed",
                         max_bytes=XLSX_MAX_UNCOMPRESSED_BYTES)


def read_xlsx(content: bytes, sheet=None, max_rows=None) -> Table:
    from openpyxl import load_workbook

    _check_zip(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl 对坏文件抛各种异常
        raise TableError('NOT_XLSX', "The file is not a valid Excel (.xlsx) file")
    try:
        sheets = list(workbook.sheetnames)
        if sheet not in (None, ''):
            if sheet not in sheets:
                raise TableError('SHEET_NOT_FOUND', f"Sheet '{sheet}' not found", sheet=sheet, sheets=sheets)
            worksheet = workbook[sheet]
        else:
            if not sheets:
                raise TableError('EMPTY', "The Excel file has no sheet")
            worksheet = workbook.worksheets[0]
        rows = []
        try:
            for number, values in enumerate(worksheet.iter_rows(values_only=True, max_col=MAX_COLUMNS), start=1):
                cells = _trim([cell_text(value) for value in values])
                if cells:
                    rows.append((number, cells))
                    if max_rows is not None and len(rows) - 1 > max_rows:
                        break
        except TableError:
            raise
        except Exception:  # noqa: BLE001
            raise TableError('NOT_XLSX', "The Excel file cannot be read")
        _limit_rows(rows, max_rows)
        return Table('xlsx', rows, sheet=worksheet.title, sheets=sheets)
    finally:
        workbook.close()


def read_table(content: bytes, fmt, encoding='auto', sheet=None, max_rows=None) -> Table:
    if fmt == 'xlsx':
        return read_xlsx(content, sheet=sheet, max_rows=max_rows)
    return read_csv(content, encoding=encoding, max_rows=max_rows)


# ------------------------------------------------------------------
# 写
# ------------------------------------------------------------------

def to_cp932_text(text: str) -> str:
    return text.translate(CP932_SUBSTITUTES)


def unencodable_chars(text, encoding) -> list:
    """text 写成 encoding 时会变成「?」的字符（去重、按出现顺序）；UTF-8 / Excel 返回 []"""
    if encoding != 'cp932' or not text:
        return []
    bad = []
    for ch in to_cp932_text(str(text)):
        if ch in bad:
            continue
        try:
            ch.encode('cp932')
        except UnicodeEncodeError:
            bad.append(ch)
    return bad


class CsvWriter:
    """逐行写 CSV（CRLF）：save() 时按编码输出（utf-8-sig 带 BOM；cp932 编码不了的写成「?」）"""

    def __init__(self, header, encoding):
        self._buffer = io.StringIO(newline='')
        self._writer = csv.writer(self._buffer, lineterminator='\r\n')
        self._encoding = encoding
        self._writer.writerow(header)

    def append(self, row):
        self._writer.writerow(['' if value is None else value for value in row])

    def save(self) -> bytes:
        text = self._buffer.getvalue()
        if self._encoding == 'cp932':
            return to_cp932_text(text).encode('cp932', errors='replace')
        return text.encode('utf-8-sig')


def write_csv(header, rows, encoding) -> bytes:
    """header + rows（单元格为文字 / 数字 / None）→ CSV 字节"""
    writer = CsvWriter(header, encoding)
    for row in rows:
        writer.append(row)
    return writer.save()


class XlsxWriter:
    """逐行写 .xlsx（write_only）：text_columns 的列按文字写（编码列保住前导零 / 长数字）"""

    def __init__(self, sheet_title, header, kinds=None):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        self._workbook = Workbook(write_only=True)
        self._sheet = self._workbook.create_sheet(title=sheet_title)
        self._kinds = list(kinds or [])
        self._sheet.freeze_panes = 'A2'
        for index, title in enumerate(header):
            letter = self._column_letter(index)
            self._sheet.column_dimensions[letter].width = max(10, min(40, len(str(title)) * 2 + 4))
        bold = Font(bold=True)
        self._append([self._cell(title, 'text', font=bold) for title in header])

    @staticmethod
    def _column_letter(index):
        from openpyxl.utils import get_column_letter
        return get_column_letter(index + 1)

    def _cell(self, value, kind, font=None):
        from openpyxl.cell import WriteOnlyCell

        if isinstance(value, str):
            value = _ILLEGAL_XML_CHARS.sub('', value)[:EXCEL_CELL_MAX_CHARS]
        cell = WriteOnlyCell(self._sheet, value=value)
        if isinstance(value, str):
            cell.data_type = 's'                  # 以 = 开头的也按文字（不当公式）
            if kind == 'text':
                cell.number_format = '@'
        elif isinstance(value, datetime):
            cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        elif isinstance(value, date):
            cell.number_format = 'yyyy-mm-dd'
        if font is not None:
            cell.font = font
        return cell

    def _append(self, cells):
        self._sheet.append(cells)

    def append(self, row):
        kinds = self._kinds
        self._append([self._cell(value, kinds[index] if index < len(kinds) else None)
                      for index, value in enumerate(row)])

    def save(self) -> bytes:
        buffer = io.BytesIO()
        self._workbook.save(buffer)
        return buffer.getvalue()


def write_xlsx(sheet_title, header, rows, kinds=None) -> bytes:
    writer = XlsxWriter(sheet_title, header, kinds)
    for row in rows:
        writer.append(row)
    return writer.save()


def content_type(fmt, encoding=None) -> str:
    if fmt == 'csv':
        return 'text/csv; charset=Shift_JIS' if encoding == 'cp932' else 'text/csv; charset=utf-8'
    return CONTENT_TYPES[fmt]


def attachment_header(filename: str) -> str:
    """Content-Disposition：ASCII 文件名 + RFC 5987 的 UTF-8 文件名（日文 / 中文文件名下载后不乱码）"""
    from urllib.parse import quote
    fallback = re.sub(r'[^A-Za-z0-9._-]+', '_', filename).strip('_') or 'download'
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quote(filename)}"
