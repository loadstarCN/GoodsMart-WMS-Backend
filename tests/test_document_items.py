"""Read-only document item imports: tenant boundaries, permissions, formats and validation."""
import io

import pytest
from openpyxl import Workbook, load_workbook
from tests.helpers import *  # noqa: F401,F403
from tests.test_inventory_availability import _h, _available
from tests.test_platform_guards import _role, _staff


@pytest.mark.parametrize('kind', ['dn', 'asn'])
def test_preview_merges_and_is_read_only(client, access_token, kind):
    with client.application.app_context():
        wid = get_warehouse().id
        before = (Goods.query.count(), Inventory.query.count())
    stock_before = _available(client, 'G001', wid)
    result = client.post(f'/{kind}/items/preview', headers=_h(access_token, wid), json={
        'text': 'code\tqty\tremark\nG001\t2\tfirst\nG001\t3\tsecond\nG002\t\t', 'has_header': True})
    assert result.status_code == 200, result.get_json()
    body = result.get_json()
    assert (body['valid_count'], body['error_count']) == (2, 0)
    first = body['items'][0]
    assert (first['lines'], first['quantity'], first['remark']) == ([2, 3], 5, 'first；second')
    assert body['items'][1]['quantity'] == 1
    assert ('available_stock_for_sale' in first['goods']) == (kind == 'dn')
    assert _available(client, 'G001', wid) == stock_before
    with client.application.app_context():
        assert (Goods.query.count(), Inventory.query.count()) == before


@pytest.mark.parametrize('kind', ['dn', 'asn'])
def test_preview_tenant_inactive_and_invalid_rows(client, access_token, kind):
    with client.application.app_context():
        wid = get_warehouse().id
        db.session.add(Goods(code='PRIVATE', name='Other tenant', company_id=2, unit='pcs', created_by=1))
        Goods.query.filter_by(code='G002').first().is_active = False
        db.session.commit()
    result = client.post(f'/{kind}/items/preview', headers=_h(access_token, wid), json={
        'text': 'PRIVATE,1\nG002,1\nG001,-1\nG001,1,ok,extra\n,2\nG001,2147483647\nG001,1'})
    assert result.status_code == 200, result.get_json()
    rows = result.get_json()['items']
    assert [r['error'] for r in rows] == ['NOT_FOUND', 'INACTIVE', 'QUANTITY_INVALID', 'EXTRA_COLUMNS', 'CODE_INVALID', 'QUANTITY_INVALID']
    assert rows[0]['goods'] is None


@pytest.mark.parametrize('kind', ['dn', 'asn'])
def test_preview_permissions_and_warehouse_required(client, access_token, kind):
    with client.application.app_context():
        warehouse = get_warehouse()
        wid = warehouse.id
        editor = _staff('item_editor', _role('item_editor', [f'{kind}_edit']), [warehouse])
        reader = _staff('item_reader', _role('item_reader', [f'{kind}_read']), [warehouse])
        editor_token, reader_token = create_access_token(identity=editor), create_access_token(identity=reader)
    endpoint = f'/{kind}/items/preview'
    assert client.post(endpoint, headers=_h(access_token), json={'text': 'G001'}).status_code == 400
    assert client.post(endpoint, headers=_h(reader_token, wid), json={'text': 'G001'}).status_code == 403
    assert client.post(endpoint, headers=_h(editor_token, wid), json={'text': 'G001'}).status_code == 200
    assert client.get(f'/{kind}/items/template', headers=_h(reader_token)).status_code == 403


def test_preview_files_and_templates(client, access_token):
    with client.application.app_context():
        wid = get_warehouse().id
    headers = _h(access_token, wid)
    response = client.post('/dn/items/preview', headers=headers, data={
        'file': (io.BytesIO('商品,数量,備考\nG001,2,日本語'.encode('shift_jis')), 'items.csv')})
    assert response.status_code == 200, response.get_json()
    assert response.get_json()['items'][0]['remark'] == '日本語'
    book = Workbook()
    book.active.title = 'Empty'
    sheet = book.create_sheet('Items')
    sheet.append(['code', 'quantity', 'remark'])
    sheet.append(['G001', 3, 'second sheet'])
    output = io.BytesIO()
    book.save(output)
    response = client.post('/asn/items/preview', headers=headers, data={
        'file': (io.BytesIO(output.getvalue()), 'items.xlsx'), 'sheet': 'Items'})
    assert response.status_code == 200, response.get_json()
    assert response.get_json()['sheets'] == ['Empty', 'Items']
    assert response.get_json()['items'][0]['quantity'] == 3
    for kind in ['dn', 'asn']:
        response = client.get(f'/{kind}/items/template?format=xlsx&lang=zh-CN', headers=headers)
        assert response.status_code == 200
        template = load_workbook(io.BytesIO(response.data))
        assert template.active['A2'].value == '0012345678901'
        assert template.active['A2'].data_type == 's'
        template.close()


@pytest.mark.parametrize('body,reason', [
    ({'text': ''}, 'EMPTY'), ({'text': 'G001', 'has_header': 'false'}, 'HEADER_INVALID'),
    ({'text': '"unterminated'}, 'NOT_CSV'),
    ({'text': '\n'.join(f'CODE{i}' for i in range(1001))}, 'TOO_MANY_CODES'),
])
def test_preview_invalid_input(client, access_token, body, reason):
    with client.application.app_context():
        wid = get_warehouse().id
    response = client.post('/dn/items/preview', headers=_h(access_token, wid), json=body)
    assert response.status_code == 400
    assert response.get_json()['details']['reason'] == reason
