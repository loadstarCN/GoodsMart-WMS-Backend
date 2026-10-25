"""商品导入二期：dry_run 预校验、按 import_token 执行（跳过错误行）、Excel / Shift_JIS / 中英日表头、错误行文件、导入记录

- dry_run：一次返回全部错误；逐行 create / update（字段级差异 from → to）/ unchanged / skip / error（行号、列、原因）；不写任何东西
- 执行带 import_token：导入其余行、跳过错误行（status partial）；文件 / 方式变了 409 16141 FILE_CHANGED，期间商品被改 DATA_CHANGED
- 编码自动识别（Shift_JIS）、日文表头、千位分隔的数字；Excel：数字 / 日期单元格、长数字编码、指定工作表
- 错误行文件：原格式 + 末列「错误」（dry_run 阶段 POST /bulk_upload/errors；执行后 GET /imports/<id>/errors），30 天后删除
- 导入记录：每次执行一条（不带 token 失败的也记）；商品列表 ?import_id= 只看这次新增 / 更新的
- 单次最多 10,000 行；模板支持 xlsx 与表头语言
"""
import csv
import io

from openpyxl import Workbook, load_workbook

from tests.helpers import *  # noqa: F401,F403
from datetime import datetime, timedelta, timezone   # helpers 里 import 的是 datetime 模块，放在其后
from warehouse.goods.models import GoodsImport, GoodsImportItem


def _h(token):
    return {'Authorization': f'Bearer {token}'}


def _post(client, token, content, filename='goods.csv', path='/goods/bulk_upload', **form):
    data = {'file': (io.BytesIO(content), filename), 'company_id': '1'}
    data.update({key: str(value).lower() if isinstance(value, bool) else str(value) for key, value in form.items()})
    return client.post(path, headers=_h(token), data=data, content_type='multipart/form-data')


def _xlsx(rows, title='商品', extra_sheet=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    if extra_sheet:
        workbook.create_sheet(extra_sheet).append(['x'])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _goods(client, code):
    with client.application.app_context():
        return Goods.query.filter_by(code=code, company_id=1).first()


JA_ROWS = [
    'コード,商品名,重量 kg,原産国,価格,割引価格,ブランド,仕入先',
    'NEW-1,新商品,0.328,JP,880,,ステラ,X',          # 2 新增
    'G001,Sample Goods,1.2,jp,,,,',                 # 3 更新：原产国（重量相同不算）
    'G002,Sample Goods,,,,,,',                      # 4 更新：停用 → 启用
    'SAME-1,Same,,,,,,',                            # 5 无变化
    'NEW-2,価格エラー,,,"4,400円",,,',               # 6 错误：不是数字
    'NEW-1,重複,,,,,,',                              # 7 错误：与第 2 行重复
    'NEW-3,,,,,,,',                                 # 8 错误：名称为空
    'NEW-4,マイナス,-0.2,,,,,',                       # 9 错误：负数
    'NEW-5,千円,,,"1,200",,,',                       # 10 新增：千位分隔
]


def _ja_file():
    return ('\r\n'.join(JA_ROWS) + '\r\n').encode('cp932')


def _prepare_existing(client):
    with client.application.app_context():
        Goods.query.filter_by(code='G002').first().is_active = False
        db.session.add(Goods(code='SAME-1', name='Same', company_id=1, unit='pcs', currency='JPY', created_by=1))
        db.session.commit()


def test_dry_run_reports_every_row_and_writes_nothing(client, access_token):
    _prepare_existing(client)
    response = _post(client, access_token, _ja_file(), overwrite='override', dry_run=True)
    body = response.get_json()
    assert response.status_code == 200, body
    assert body['dry_run'] is True and len(body['import_token']) == 64
    assert body['file']['format'] == 'csv' and body['file']['encoding'] == 'cp932'
    assert body['file']['encoding_detected'] is True and body['file']['rows'] == 9
    assert body['header_lang'] == 'ja-JP'
    assert body['columns'] == ['code', 'name', 'weight', 'origin_country', 'price', 'discount_price', 'brand']
    assert body['ignored_columns'] == ['仕入先']
    assert body['summary'] == {'rows': 9, 'create': 2, 'update': 2, 'unchanged': 1, 'skip': 0, 'error': 4}

    rows = {row['row']: row for row in body['rows']}
    assert rows[2]['action'] == 'create' and rows[2]['values']['brand'] == 'ステラ'
    assert (rows[2]['values']['price'], rows[2]['values']['currency'], rows[2]['values']['unit']) == (880, 'JPY', 'pcs')
    assert rows[2]['spec_missing'] == ['dimensions']
    assert rows[3]['action'] == 'update' and rows[3]['changes'] == [
        {'field': 'origin_country', 'from': None, 'to': 'JP'}]
    assert rows[4]['changes'] == [{'field': 'is_active', 'from': False, 'to': True}]
    assert rows[5]['action'] == 'unchanged' and 'changes' not in rows[5]
    assert rows[10]['action'] == 'create' and rows[10]['values']['price'] == 1200

    error = rows[6]['errors'][0]
    assert (error['field'], error['column'], error['code'], error['reason'], error['value']) == (
        'price', '価格', 10012, 'INVALID_NUMBER', '4,400円')
    assert rows[7]['errors'][0]['reason'] == 'DUPLICATE' and rows[7]['errors'][0]['first_row'] == 2
    assert rows[7]['errors'][0]['code'] == 16101
    assert (rows[8]['errors'][0]['reason'], rows[8]['errors'][0]['field'], rows[8]['errors'][0]['code']) == (
        'REQUIRED', 'name', 10006)
    assert (rows[9]['errors'][0]['reason'], rows[9]['errors'][0]['code']) == ('NEGATIVE', 16100)

    # 什么都没写
    assert _goods(client, 'NEW-1') is None and _goods(client, 'G001').origin_country is None
    assert _goods(client, 'G002').is_active is False
    with client.application.app_context():
        assert GoodsImport.query.count() == 0

    # 只新增：已存在的编码 → skip
    body = _post(client, access_token, _ja_file(), dry_run=True).get_json()
    assert body['summary']['skip'] == 3 and body['summary']['update'] == 0


def test_execute_with_token_skips_error_rows_and_records_the_import(client, access_token):
    _prepare_existing(client)
    token = _post(client, access_token, _ja_file(), overwrite='override', dry_run=True).get_json()['import_token']
    response = _post(client, access_token, _ja_file(), overwrite='override', import_token=token)
    body = response.get_json()
    assert response.status_code == 200, body
    assert body['status'] == 'partial' and body['has_error_file'] is True
    assert (body['created'], body['updated'], body['unchanged'], body['skipped'], body['errors']) == (2, 2, 1, 0, 4)
    assert _goods(client, 'NEW-1').brand == 'ステラ' and float(_goods(client, 'NEW-5').price) == 1200
    assert _goods(client, 'G001').origin_country == 'JP' and _goods(client, 'G002').is_active is True
    assert _goods(client, 'NEW-2') is None and _goods(client, 'NEW-4') is None

    import_id = body['import_id']
    with client.application.app_context():
        items = {(item.goods_id, item.action) for item in GoodsImportItem.query.filter_by(import_id=import_id)}
        assert len(items) == 4

    # 商品列表按导入筛选
    listed = client.get(f'/goods/?import_id={import_id}&per_page=50', headers=_h(access_token)).get_json()
    assert {item['code'] for item in listed['items']} == {'NEW-1', 'NEW-5', 'G001', 'G002'}
    listed = client.get(f'/goods/?import_id={import_id}&import_action=create', headers=_h(access_token)).get_json()
    assert {item['code'] for item in listed['items']} == {'NEW-1', 'NEW-5'}
    listed = client.get(f'/goods/?import_id={import_id}&import_action=create&spec_missing=weight',
                        headers=_h(access_token)).get_json()
    assert [item['code'] for item in listed['items']] == ['NEW-5']

    # 导入记录
    records = client.get('/goods/imports', headers=_h(access_token)).get_json()
    assert records['total'] == 1
    record = records['items'][0]
    assert (record['id'], record['status'], record['mode'], record['encoding']) == (import_id, 'partial', 'override',
                                                                                    'cp932')
    assert (record['created'], record['updated'], record['unchanged'], record['errors']) == (2, 2, 1, 4)
    assert record['has_original_file'] and record['has_error_file'] and record['created_by_user']['user_name'] == 'admin'
    detail = client.get(f'/goods/imports/{import_id}', headers=_h(access_token)).get_json()
    assert [row['row'] for row in detail['error_rows']] == [6, 7, 8, 9]

    original = client.get(f'/goods/imports/{import_id}/file', headers=_h(access_token))
    assert original.status_code == 200 and original.data == _ja_file()
    errors = client.get(f'/goods/imports/{import_id}/errors', headers=_h(access_token))
    assert errors.status_code == 200 and 'Shift_JIS' in errors.headers['Content-Type']
    lines = list(csv.reader(io.StringIO(errors.data.decode('cp932'), newline='')))
    assert lines[0] == JA_ROWS[0].split(',') + ['エラー']
    assert [line[0] for line in lines[1:]] == ['NEW-2', 'NEW-1', 'NEW-3', 'NEW-4']
    assert lines[1][4] == '4,400円' and lines[1][-1] == '価格: 数値ではありません：4,400円'
    assert lines[2][-1] == 'コード: 2 行目と重複しています'

    # 30 天后文件删除：记录还在，下载 409 16142
    from warehouse.goods.importer import purge_expired_files
    with client.application.app_context():
        assert purge_expired_files(datetime.now(timezone.utc) + timedelta(days=31)) == 1
    response = client.get(f'/goods/imports/{import_id}/file', headers=_h(access_token))
    assert (response.status_code, response.get_json()['code']) == (409, 16142)
    assert client.get(f'/goods/imports/{import_id}', headers=_h(access_token)).get_json()['has_original_file'] is False


def test_token_mismatch_and_invalid_parameters(client, access_token):
    _prepare_existing(client)
    token = _post(client, access_token, _ja_file(), overwrite='override', dry_run=True).get_json()['import_token']

    # 方式变了 → FILE_CHANGED
    response = _post(client, access_token, _ja_file(), overwrite='append', import_token=token)
    assert (response.status_code, response.get_json()['code']) == (409, 16141)
    assert response.get_json()['details'] == {'reason': 'FILE_CHANGED'}
    # 期间商品被改了（G001 已有原产国 → 这一行的结果不同）→ DATA_CHANGED，什么都没导入
    with client.application.app_context():
        Goods.query.filter_by(code='G001').first().origin_country = 'CN'
        db.session.commit()
    response = _post(client, access_token, _ja_file(), overwrite='override', import_token=token)
    assert (response.status_code, response.get_json()['details']) == (409, {'reason': 'DATA_CHANGED'})
    assert _goods(client, 'NEW-1') is None
    # token 格式不对 / 编码 / 语言参数不对 → 400 10019
    response = _post(client, access_token, _ja_file(), overwrite='override', import_token='abc')
    assert (response.get_json()['code'], response.get_json()['field']) == (10019, 'import_token')
    response = _post(client, access_token, _ja_file(), encoding='latin-1', dry_run=True)
    assert (response.get_json()['code'], response.get_json()['field']) == (10019, 'encoding')
    response = _post(client, access_token, _ja_file(), lang='fr', dry_run=True)
    assert (response.get_json()['code'], response.get_json()['field']) == (10019, 'lang')
    # 手动指定 Shift_JIS
    body = _post(client, access_token, _ja_file(), encoding='shift_jis', dry_run=True).get_json()
    assert body['file']['encoding'] == 'cp932' and body['file']['encoding_detected'] is False
    with client.application.app_context():
        assert GoodsImport.query.count() == 0


def test_without_token_errors_reject_the_whole_file_and_are_recorded(client, access_token):
    response = _post(client, access_token, 'code,name,price\nOK-1,Fine,100\nBAD-1,Bad,abc\n'.encode('utf-8'))
    assert (response.status_code, response.get_json()['code']) == (400, 10012)
    assert _goods(client, 'OK-1') is None
    response = _post(client, access_token, 'code,title\nX1,One\n'.encode('utf-8'), filename='items.csv')
    assert (response.status_code, response.get_json()['code']) == (400, 10005)

    records = client.get('/goods/imports', headers=_h(access_token)).get_json()['items']
    assert [(r['status'], r['failure']['code'], r['errors'], r['created']) for r in records] == [
        ('failed', 10005, 0, 0), ('failed', 10012, 1, 0)]
    assert records[0]['file_name'] == 'items.csv' and records[1]['has_error_file'] is True
    lines = client.get(f"/goods/imports/{records[1]['id']}/errors", headers=_h(access_token)).data.decode('utf-8-sig')
    assert lines.splitlines() == ['code,name,price,Error', 'BAD-1,Bad,abc,price: not a valid number: abc']


def test_error_file_before_import(client, access_token):
    # Shift_JIS 的文件：要中文也写日文（Shift_JIS 写不下简体字）
    response = _post(client, access_token, _ja_file(), path='/goods/bulk_upload/errors', lang='zh-CN')
    assert response.status_code == 200 and 'filename*=UTF-8' in response.headers['Content-Disposition']
    lines = list(csv.reader(io.StringIO(response.data.decode('cp932'), newline='')))
    assert lines[0][-1] == 'エラー' and len(lines) == 5
    # UTF-8 的文件按 lang 写
    response = _post(client, access_token, _ja_file().decode('cp932').encode('utf-8'),
                     path='/goods/bulk_upload/errors', lang='zh-CN')
    lines = list(csv.reader(io.StringIO(response.data.decode('utf-8-sig'), newline='')))
    assert lines[0][-1] == '错误' and lines[1][-1] == '価格: 不是有效数字：4,400円'
    # 没有错误行 → 204
    response = _post(client, access_token, 'code,name\nOK-9,Fine\n'.encode('utf-8'), path='/goods/bulk_upload/errors')
    assert response.status_code == 204
    # Excel：错误行文件也是 Excel，末列「错误」
    content = _xlsx([['编码', '名称', '价格'], ['E-1', '好', 10], ['E-2', '坏', 'abc']])
    response = _post(client, access_token, content, filename='商品.xlsx', path='/goods/bulk_upload/errors')
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert values == [['编码', '名称', '价格', '错误'], ['E-2', '坏', 'abc', '价格: 不是有效数字：abc']]
    assert sheet.title == '商品'


def test_xlsx_import_with_typed_cells_and_sheet(client, access_token):
    content = _xlsx([
        ['编码', '名称', '重量 kg', '长 mm', '宽 mm', '高 mm', '价格', '有效期至', '标签'],
        [4900000000137, 'Excel 商品', 0.5, 120, 80.0, 45, 1760, datetime(2026, 12, 31), 'a,b'],
        [None, None, None, None, None, None, None, None, None],
        ['XL-2', 'Second', '0.25', '10', '10', '10', None, '2026/1/5', None],
    ], extra_sheet='Other')
    body = _post(client, access_token, content, filename='goods.xlsx', dry_run=True).get_json()
    assert body['file']['format'] == 'xlsx' and body['file']['sheet'] == '商品'
    assert body['file']['sheets'] == ['商品', 'Other'] and body['header_lang'] == 'zh-CN'
    assert [row['row'] for row in body['rows']] == [2, 4]          # 空行跳过，行号照算
    response = _post(client, access_token, content, filename='goods.xlsx')
    assert response.status_code == 200 and response.get_json()['created'] == 2, response.get_json()
    goods = _goods(client, '4900000000137')
    assert goods.name == 'Excel 商品' and float(goods.weight) == 0.5 and (goods.length, goods.width) == (120, 80)
    assert goods.expiration_date.isoformat() == '2026-12-31' and goods.tags == 'a,b' and float(goods.price) == 1760
    assert _goods(client, 'XL-2').expiration_date.isoformat() == '2026-01-05'

    response = _post(client, access_token, content, filename='goods.xlsx', sheet='Nope', dry_run=True)
    body = response.get_json()
    assert (body['code'], body['field'], body['details']['sheets']) == (10019, 'sheet', ['商品', 'Other'])
    response = _post(client, access_token, b'not a zip', filename='goods.xlsx', dry_run=True)
    assert (response.get_json()['code'], response.get_json()['details']['reason']) == (10016, 'NOT_XLSX')


def test_row_limit_and_template_formats(client, access_token):
    rows = 'code,name\n' + ''.join(f'L{i},Goods {i}\n' for i in range(10001))
    response = _post(client, access_token, rows.encode('utf-8'), dry_run=True)
    body = response.get_json()
    assert (response.status_code, body['code'], body['details']['max_rows']) == (400, 10018, 10000)

    response = client.get('/goods/bulk_upload/template?format=xlsx&lang=ja-JP', headers=_h(access_token))
    assert response.status_code == 200 and response.headers['Content-Type'].startswith(
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    sheet = load_workbook(io.BytesIO(response.data)).active
    header, sample = [[cell.value for cell in row] for row in sheet.iter_rows()]
    assert sheet.title == '商品' and header[:4] == ['コード', '商品名', '単位', '重量 kg'] and len(header) == 20
    assert sample[0] == 'SAMPLE-001' and sample[3] == 0.328 and sample[4] == 140
    # Excel 模板原样导入
    body = _post(client, access_token, response.data, filename='Goods_template.xlsx').get_json()
    assert body['created'] == 1 and body['ignored_columns'] == []
    assert float(_goods(client, 'SAMPLE-001').weight) == 0.328

    text = client.get('/goods/bulk_upload/template?lang=zh-CN', headers=_h(access_token)).data.decode('utf-8-sig')
    assert text.splitlines()[0].startswith('编码,名称,单位,重量 kg,长 mm')
    response = client.get('/goods/bulk_upload/template?lang=xx', headers=_h(access_token))
    assert (response.status_code, response.get_json()['code']) == (400, 10019)
