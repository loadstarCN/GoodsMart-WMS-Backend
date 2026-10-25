"""商品导出（/goods/exports）

- 列清单与预设；筛选沿用商品列表（goods_codes、is_active=all、has_stock、创建 / 更新时间……）
- ≤ 1,000 件同步生成（201 completed）：CSV UTF-8（BOM）/ Shift_JIS（〜 换成 ～，写不下的「?」，预检统计、简体字表头换日文）/ Excel；
  表头按语言；「导入模板格式」导出的文件原样导回是「无变化」
- 库存列按库存范围合计；限仓库的员工只看自己仓库的、或不含库存的导出
- 超过阈值走后台：202 pending → 定时任务生成（进度）→ completed；失败重试、中断恢复、7 天后删文件（expired）、按原条件重新导出
- 参数校验 400 10021、件数上限 400 10022、不能下载 409 16142、不能重试 409 16143
"""
import csv
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import lazyload

from tests.helpers import *  # noqa: F401,F403
from tests.test_platform_guards import _role, _staff
from datetime import datetime, timedelta, timezone   # helpers 里 import 的是 datetime 模块，放在其后
from warehouse.goods import exporter
from warehouse.goods.columns import GOODS_IMPORT_COLUMNS, header_label
from warehouse.goods.models import GoodsExport


def _h(token, warehouse_id=None):
    headers = {'Authorization': f'Bearer {token}'}
    if warehouse_id is not None:
        headers['X-WAREHOUSE-ID'] = str(warehouse_id)
    return headers


def _export(client, token, warehouse_id=None, **body):
    return client.post('/goods/exports', headers=_h(token, warehouse_id), json=body)


def _download(client, token, export_id, warehouse_id=None):
    return client.get(f'/goods/exports/{export_id}/download', headers=_h(token, warehouse_id))


def _csv(response, encoding='utf-8-sig'):
    return list(csv.reader(io.StringIO(response.data.decode(encoding), newline='')))


def _inventory(client, code, warehouse_id=None):
    with client.application.app_context():
        goods = Goods.query.filter_by(code=code).first()
        query = Inventory.query.options(lazyload('*')).filter_by(goods_id=goods.id)
        if warehouse_id:
            query = query.filter_by(warehouse_id=warehouse_id)
        rows = query.all()
        return {'total': sum(r.total_stock for r in rows), 'available': sum(r.available_stock_for_sale for r in rows),
                'pending': sum(r.sorted_stock for r in rows), 'damaged': sum(r.damage_stock for r in rows)}


def test_columns_catalog(client, access_token):
    body = client.get('/goods/exports/columns', headers=_h(access_token)).get_json()
    assert body['presets']['import_template'] == list(GOODS_IMPORT_COLUMNS)
    assert len(body['presets']['all']) == 28 and body['langs'] == ['zh-CN', 'en-US', 'ja-JP']
    code = next(column for column in body['columns'] if column['key'] == 'code')
    assert code['labels'] == {'zh-CN': '编码', 'en-US': 'Code', 'ja-JP': 'コード'} and code['required']
    stock = next(column for column in body['columns'] if column['key'] == 'stock_available')
    assert stock['group'] == 'stock' and stock['importable'] is False


def test_sync_csv_export_round_trips_through_the_import(client, access_token):
    response = _export(client, access_token, filters={'goods_codes': ['G001', 'G002']}, lang='zh-CN')
    body = response.get_json()
    assert response.status_code == 201, body
    assert (body['status'], body['background'], body['preset'], body['total'], body['row_count']) == (
        'completed', False, 'import_template', 2, 2)
    assert (body['file_format'], body['encoding'], body['progress'], body['downloadable']) == ('csv', 'utf-8-sig',
                                                                                                100, True)
    assert body['expires_at'] and body['created_by_user']['user_name'] == 'admin'

    response = _download(client, access_token, body['id'])
    assert response.status_code == 200 and response.data.startswith(b'\xef\xbb\xbf')
    assert body['file_name'] in response.headers['Content-Disposition']
    rows = _csv(response)
    assert rows[0] == [header_label(key, 'zh-CN') for key in GOODS_IMPORT_COLUMNS]
    assert [row[0] for row in rows[1:]] == ['G002', 'G001']        # 与列表同序（新的在前）
    assert rows[2][3] == '1.2' and rows[2][4] == '10'               # 重量去掉多余的 0、尺寸整数

    # 导入模板格式：原样导回 → 全部「无变化」
    preview = client.post('/goods/bulk_upload', headers=_h(access_token), data={
        'file': (io.BytesIO(response.data), 'export.csv'), 'company_id': '1', 'overwrite': 'override',
        'dry_run': 'true'}, content_type='multipart/form-data').get_json()
    assert preview['summary']['unchanged'] == 2 and preview['ignored_columns'] == [], preview

    # 列表 / 详情
    listed = client.get('/goods/exports', headers=_h(access_token)).get_json()
    assert listed['total'] == 1 and listed['items'][0]['id'] == body['id']
    assert client.get(f"/goods/exports/{body['id']}", headers=_h(access_token)).get_json()['status'] == 'completed'


def test_shift_jis_preview_counts_unencodable_goods(client, access_token):
    with client.application.app_context():
        Goods.query.filter_by(code='G001').first().name = 'ハート♡キーホルダー'
        Goods.query.filter_by(code='G002').first().name = 'まほろば〜学園'
        db.session.commit()
    body = {'filters': {'goods_codes': 'G001,G002'}, 'encoding': 'shift_jis', 'lang': 'zh-CN'}
    preview = client.post('/goods/exports/preview', headers=_h(access_token), json=body).get_json()
    assert preview['total'] == 2 and preview['background'] is False
    unencodable = preview['unencodable']
    assert unencodable['count'] == 1 and unencodable['characters'] == ['♡']
    assert [(item['code'], item['columns']) for item in unencodable['items']] == [('G001', ['name'])]
    # 简体字表头 Shift_JIS 写不下 → 换成日文表头（导入照样认）
    headers = {column['key']: column['header'] for column in preview['columns']}
    assert headers['code'] == 'コード' and headers['name'] == '名称'

    created = _export(client, access_token, **body).get_json()
    assert created['unencodable_count'] == 1 and created['encoding'] == 'cp932'
    response = _download(client, access_token, created['id'])
    assert 'Shift_JIS' in response.headers['Content-Type']
    rows = {row[0]: row for row in _csv(response, 'cp932')[1:]}
    assert rows['G001'][1] == 'ハート?キーホルダー' and rows['G002'][1] == 'まほろば～学園'
    # UTF-8 不预估
    preview = client.post('/goods/exports/preview', headers=_h(access_token),
                          json={**body, 'encoding': 'utf-8'}).get_json()
    assert preview['unencodable'] is None


def test_xlsx_all_columns_with_stock_scope_and_filters(client, access_company_admin_token):
    token = access_company_admin_token
    with client.application.app_context():
        warehouse_a, warehouse_b = get_warehouse().id, Warehouse.query.filter_by(name='Warehouse Test').first().id
        g1 = Goods.query.filter_by(code='G001').first()
        add_stock(g1.id, warehouse_b, onhand=7)
        db.session.add(Goods(code='EMPTY-1', name='No stock', company_id=1, unit='pcs', created_by=1,
                             is_active=False))
        db.session.commit()

    response = _export(client, token, preset='all', file_format='xlsx', lang='en-US',
                       filters={'goods_codes': ['G001', 'G002', 'EMPTY-1'], 'is_active': 'all'})
    body = response.get_json()
    assert response.status_code == 201 and body['preset'] == 'all' and body['encoding'] is None, body
    sheet = load_workbook(io.BytesIO(_download(client, token, body['id']).data)).active
    rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
    header = rows[0]
    assert sheet.title == 'Goods' and header[:2] == ['Code', 'Name'] and header[-4:] == [
        'Status', 'Created at', 'Updated at', 'Created by']
    by_code = {row[0]: dict(zip(header, row)) for row in rows[1:]}
    assert set(by_code) == {'G001', 'G002', 'EMPTY-1'}
    expected = _inventory(client, 'G001')           # 公司管理员没选仓库：全部仓库合计
    assert (by_code['G001']['On hand'], by_code['G001']['Available']) == (expected['total'], expected['available'])
    assert by_code['EMPTY-1']['On hand'] == 0 and by_code['EMPTY-1']['Status'] == 'Inactive'
    assert by_code['G001']['Status'] == 'Active' and by_code['G001']['Created by'] == 'admin'
    assert isinstance(by_code['G001']['Weight kg'], (int, float)) and isinstance(by_code['G001']['Created at'],
                                                                                 datetime)

    # 只算仓库 A
    body = _export(client, token, columns=['code', 'stock_on_hand'], lang='ja-JP',
                   filters={'goods_codes': ['G001'], 'warehouse_id': warehouse_a}).get_json()
    assert body['warehouse_id'] == warehouse_a and body['preset'] == 'custom'
    rows = _csv(_download(client, token, body['id']))
    assert rows == [['コード', '在庫'], ['G001', str(_inventory(client, 'G001', warehouse_a)['total'])]]

    # 有 / 无库存（列表同样支持）
    body = _export(client, token, columns=['code'], filters={'has_stock': False, 'is_active': 'all',
                                                             'goods_codes': ['G001', 'EMPTY-1']}).get_json()
    assert [row[0] for row in _csv(_download(client, token, body['id']))[1:]] == ['EMPTY-1']
    listed = client.get('/goods/?has_stock=true&is_active=all&goods_codes=G001,EMPTY-1',
                        headers=_h(token)).get_json()
    assert [item['code'] for item in listed['items']] == ['G001']


def test_created_and_updated_time_filters(client, access_token):
    with client.application.app_context():
        g1, g2 = Goods.query.filter_by(code='G001').first(), Goods.query.filter_by(code='G002').first()
        g1.created_at = datetime(2026, 9, 30, 10, 0, tzinfo=timezone.utc)    # UTC / 东京都还是 9/30
        g2.created_at = datetime(2026, 10, 10, 3, 0, tzinfo=timezone.utc)
        db.session.commit()
    listed = client.get('/goods/?created_from=2026-10-01&created_to=2026-10-10', headers=_h(access_token)).get_json()
    assert [item['code'] for item in listed['items']] == ['G002']
    listed = client.get('/goods/?created_to=2026-09-30T23:00:00Z', headers=_h(access_token)).get_json()
    assert [item['code'] for item in listed['items']] == ['G001']
    body = _export(client, access_token, columns=['code'], filters={'created_from': '2026-10-01'}).get_json()
    assert body['filters'] == {'created_from': '2026-10-01'}
    assert 'G001' not in [row[0] for row in _csv(_download(client, access_token, body['id']))]
    response = client.get('/goods/?created_from=yesterday', headers=_h(access_token))
    assert response.status_code == 400


def test_background_export_progress_retry_expiry_and_rerun(client, access_token, monkeypatch):
    monkeypatch.setattr(exporter, 'EXPORT_SYNC_LIMIT', 1)
    monkeypatch.setattr(exporter, 'EXPORT_BATCH_SIZE', 1)
    response = _export(client, access_token, filters={'goods_codes': ['G001', 'G002']}, file_format='xlsx')
    body = response.get_json()
    assert response.status_code == 202 and (body['status'], body['background'], body['progress']) == (
        'pending', True, 0), body
    response = _download(client, access_token, body['id'])
    assert (response.status_code, response.get_json()['code'], response.get_json()['details']) == (
        409, 16142, {'status': 'pending'})

    with client.application.app_context():
        result = exporter.run_goods_file_jobs(time_budget=10)
    assert result['generated'] == 1
    detail = client.get(f"/goods/exports/{body['id']}", headers=_h(access_token)).get_json()
    assert (detail['status'], detail['processed'], detail['row_count'], detail['attempts']) == ('completed', 2, 2, 1)
    assert _download(client, access_token, body['id']).status_code == 200

    # 只有 failed 的能重试
    response = client.post(f"/goods/exports/{body['id']}/retry", headers=_h(access_token))
    assert (response.status_code, response.get_json()['code']) == (409, 16143)

    # 生成失败 → failed（带原因）→ 重试成功
    failing = _export(client, access_token, filters={'goods_codes': ['G001', 'G002']}).get_json()
    original = exporter._generate
    monkeypatch.setattr(exporter, '_generate', lambda export: (_ for _ in ()).throw(RuntimeError('disk full')))
    with client.application.app_context():
        exporter.run_goods_file_jobs(time_budget=10)
    detail = client.get(f"/goods/exports/{failing['id']}", headers=_h(access_token)).get_json()
    assert detail['status'] == 'failed' and detail['error'] == 'RuntimeError: disk full'
    monkeypatch.setattr(exporter, '_generate', original)
    retried = client.post(f"/goods/exports/{failing['id']}/retry", headers=_h(access_token))
    assert retried.status_code == 202 and retried.get_json()['status'] == 'pending'
    with client.application.app_context():
        exporter.run_goods_file_jobs(time_budget=10)
        assert db.session.get(GoodsExport, failing['id']).attempts == 2
    assert client.get(f"/goods/exports/{failing['id']}", headers=_h(access_token)).get_json()['status'] == 'completed'

    # 进程中断（heartbeat 过期）→ 重新排队；超过 3 次 → failed
    with client.application.app_context():
        export = db.session.get(GoodsExport, failing['id'])
        export.status, export.heartbeat_at = 'running', datetime.now(timezone.utc) - timedelta(minutes=30)
        db.session.commit()
        assert exporter.recover_stale() == {'requeued': 1, 'failed': 0}
        export = db.session.get(GoodsExport, failing['id'])
        export.status, export.attempts, export.heartbeat_at = 'running', 3, None
        db.session.commit()
        assert exporter.recover_stale() == {'requeued': 0, 'failed': 1}

    # 7 天后删文件 → expired，下载 409；按原条件重新导出 = 新记录
    with client.application.app_context():
        assert exporter.purge_expired(datetime.now(timezone.utc) + timedelta(days=8)) == 1
    response = _download(client, access_token, body['id'])
    assert (response.status_code, response.get_json()['details']) == (409, {'status': 'expired'})
    rerun = client.post(f"/goods/exports/{body['id']}/rerun", headers=_h(access_token))
    data = rerun.get_json()
    assert rerun.status_code == 202 and data['source_export_id'] == body['id'] and data['file_format'] == 'xlsx'
    assert data['id'] != body['id'] and data['filters'] == {'goods_codes': ['G001', 'G002']}


def test_restricted_staff_only_see_their_warehouse_stock_exports(client, access_company_admin_token):
    with client.application.app_context():
        warehouse_a = get_warehouse()
        staff = _staff('goods_reader', _role('goods_reader', ['goods_read']), [warehouse_a])
        staff_token = create_access_token(identity=staff)
        warehouse_a_id = warehouse_a.id
    # 公司管理员：全部仓库的库存 / 不含库存 / 仓库 A 的库存
    all_stock = _export(client, access_company_admin_token, columns=['code', 'stock_on_hand']).get_json()
    plain = _export(client, access_company_admin_token, columns=['code', 'name']).get_json()
    stock_a = _export(client, access_company_admin_token, warehouse_id=warehouse_a_id,
                      columns=['code', 'stock_available']).get_json()
    listed = client.get('/goods/exports', headers=_h(staff_token, warehouse_a_id)).get_json()
    assert {item['id'] for item in listed['items']} == {plain['id'], stock_a['id']}
    response = client.get(f"/goods/exports/{all_stock['id']}", headers=_h(staff_token, warehouse_a_id))
    assert response.status_code == 404
    # 员工自己导出：库存范围就是当前仓库；指定别的仓库 → 403
    own = _export(client, staff_token, warehouse_a_id, columns=['code', 'stock_on_hand']).get_json()
    assert own['warehouse_id'] == warehouse_a_id
    with client.application.app_context():
        warehouse_b_id = Warehouse.query.filter_by(name='Warehouse Test').first().id
    response = _export(client, staff_token, warehouse_a_id, columns=['code'],
                       filters={'warehouse_id': warehouse_b_id})
    assert response.status_code == 403


@pytest.mark.parametrize('body, field', [
    ({'columns': ['code', 'colour']}, 'columns'),
    ({'columns': ['name']}, 'columns'),
    ({'columns': ['code', 'code']}, 'columns'),
    ({'preset': 'everything'}, 'preset'),
    ({'file_format': 'pdf'}, 'file_format'),
    ({'encoding': 'latin-1'}, 'encoding'),
    ({'lang': 'fr-FR'}, 'lang'),
    ({'filters': {'colour': 'red'}}, 'filters.colour'),
    ({'filters': {'has_stock': 'maybe'}}, 'filters.has_stock'),
    ({'filters': []}, 'filters'),
])
def test_invalid_export_requests(client, access_token, body, field):
    response = _export(client, access_token, **body)
    assert (response.status_code, response.get_json()['code'], response.get_json()['field']) == (400, 10021, field)


def test_export_row_limit(client, access_token, monkeypatch):
    monkeypatch.setattr(exporter, 'EXPORT_MAX_ROWS', 1)
    response = _export(client, access_token)
    body = response.get_json()
    assert (response.status_code, body['code'], body['details']['max_rows']) == (400, 10022, 1)
    preview = client.post('/goods/exports/preview', headers=_h(access_token), json={}).get_json()
    assert preview['too_large'] is True
    with client.application.app_context():
        assert GoodsExport.query.count() == 0
